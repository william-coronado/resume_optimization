from docx import Document

def docx_to_markdown(docx_file, markdown_file):
    """
    Converts a .docx file to a markdown file.

    :param docx_file: Path to the input .docx file
    :param markdown_file: Path to the output markdown file
    """
    # Load the .docx file
    document = Document(docx_file)
    
    # Open the markdown file for writing
    with open(markdown_file, 'w', encoding='utf-8') as md_file:
        for paragraph in document.paragraphs:
            # Check for heading styles
            if paragraph.style.name.startswith('Heading'):
                level = int(paragraph.style.name.split()[-1])  # Extract heading level
                md_file.write('#' * level + ' ' + paragraph.text + '\n\n')
            else:
                # Process runs for inline styles
                md_line = ""
                for run in paragraph.runs:
                    text = run.text
                    if run.bold:
                        text = f"**{text}**"
                    if run.italic:
                        text = f"*{text}*"
                    md_line += text
                md_file.write(md_line + '\n\n')

        # Process tables (if any)
        for table in document.tables:
            for row in table.rows:
                row_data = "| " + " | ".join(cell.text.strip() for cell in row.cells) + " |\n"
                md_file.write(row_data)
            md_file.write('\n')

def markdown_to_docx(markdown_file, docx_file):
    """
    Converts a markdown file to a .docx file.

    :param markdown_file: Path to the input markdown file
    :param docx_file: Path to the output .docx file
    """
    # Create a new .docx document
    document = Document()

    # Open the markdown file for reading
    with open(markdown_file, 'r', encoding='utf-8') as md_file:
        for line in md_file:
            line = line.strip()
            if line.startswith("# "):  # H1
                document.add_heading(line[2:], level=1)
            elif line.startswith("## "):  # H2
                document.add_heading(line[3:], level=2)
            elif line.startswith("### "):  # H3
                document.add_heading(line[4:], level=3)
            elif line.startswith("#### "):  # H4
                document.add_heading(line[5:], level=4)
            elif line.startswith("|") and line.endswith("|"):  # Table row
                # Parse table rows
                cells = [cell.strip() for cell in line.split("|")[1:-1]]
                if not hasattr(markdown_to_docx, "table"):
                    markdown_to_docx.table = document.add_table(rows=1, cols=len(cells))
                    hdr_cells = markdown_to_docx.table.rows[0].cells
                    for i, cell in enumerate(cells):
                        hdr_cells[i].text = cell
                else:
                    row_cells = markdown_to_docx.table.add_row().cells
                    for i, cell in enumerate(cells):
                        row_cells[i].text = cell
            else:
                # Process inline styles for bold and italic
                paragraph = document.add_paragraph()
                tokens = line.split("**")  # Split by bold markers
                for i, token in enumerate(tokens):
                    if i % 2 == 0:  # Not bold
                        sub_tokens = token.split("*")  # Split by italic markers
                        for j, sub_token in enumerate(sub_tokens):
                            if j % 2 == 0:  # Not italic
                                paragraph.add_run(sub_token)
                            else:  # Italic
                                paragraph.add_run(sub_token).italic = True
                    else:  # Bold
                        sub_tokens = token.split("*")  # Split by italic markers
                        for j, sub_token in enumerate(sub_tokens):
                            if j % 2 == 0:  # Not italic
                                paragraph.add_run(sub_token).bold = True
                            else:  # Bold and italic
                                run = paragraph.add_run(sub_token)
                                run.bold = True
                                run.italic = True

    # Save the .docx file
    document.save(docx_file)

